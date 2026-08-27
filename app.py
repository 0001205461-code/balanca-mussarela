import io
import os
import re
from datetime import datetime

import pandas as pd
import plotly.express as px
import requests
import streamlit as st


# =========================================================
# CONFIGURAÇÃO
# =========================================================

st.set_page_config(
    page_title="AMIRA | Monitoramento de Produção",
    page_icon="⚡",
    layout="wide",
    initial_sidebar_state="expanded"
)

URL_SCRIPT = "https://script.google.com/macros/s/AKfycbwQ8IIVRIDsx8-CdJeKw6LUr4rBOFGX0jb42augc8v89TZVNWy0O8mlBAK23O2Tjymmaw/exec"
TZ = "America/Sao_Paulo"


# =========================================================
# CSS
# =========================================================

st.markdown("""
<style>

@import url('https://fonts.googleapis.com/css2?family=Inter:wght@400;600;700;800&family=Orbitron:wght@600;700;800&display=swap');

:root{
 --bg:#05070d;
 --panel:#0b0f1a;
 --blue:#2f80ff;
 --cyan:#00d2ff;
 --purple:#8a35ff;
 --green:#20e889;
 --text:#f4f7ff;
 --muted:#8f9bb2;
}

html,body,[class*="css"]{
 font-family:Inter,sans-serif;
}

.stApp{
 background:
 radial-gradient(circle at 68% 12%,rgba(54,83,180,.13),transparent 28%),
 radial-gradient(circle at 96% 76%,rgba(138,53,255,.12),transparent 31%),
 var(--bg);
 color:var(--text);
}

header[data-testid="stHeader"]{
 background:transparent!important;
}

[data-testid="stDecoration"]{
 display:none!important;
}

.block-container{
 padding:1.2rem 1.6rem 2rem;
 max-width:1700px;
}


/* =========================================================
   SIDEBAR
   ========================================================= */

section[data-testid="stSidebar"]{
 background:linear-gradient(180deg,#070a12,#080b14)!important;
 border-right:1px solid rgba(76,122,255,.18);
}

section[data-testid="stSidebar"] .stButton{
 width:100%;
 margin-bottom:8px;
}

section[data-testid="stSidebar"] .stButton button{
 width:100%;
 min-height:52px;
 border-radius:10px!important;
 background:#0d1422!important;
 color:#dce6fa!important;
 border:1px solid rgba(76,122,255,.25)!important;
 text-align:left!important;
 font-weight:700;
}

section[data-testid="stSidebar"] .stButton button[kind="primary"]{
 background:linear-gradient(90deg,#1574ff,#7031eb)!important;
 color:#fff!important;
}


/* =========================================================
   MARCA
   ========================================================= */

.brand{
 text-align:center;
 font-family:Orbitron;
 font-size:1.1rem;
 font-weight:800;
}

.brand span,
.top-title span{
 color:var(--cyan);
}

.caption{
 text-align:center;
 color:var(--muted);
 font-size:.75rem;
}


/* =========================================================
   CABEÇALHO
   ========================================================= */

.top-title{
 font-size:2.35rem;
 font-weight:800;
}

.subtitle{
 color:#a0aabd;
 font-size:1rem;
 margin-top:4px;
}

.line{
 height:1px;
 background:linear-gradient(
  90deg,
  #3c78ff,
  #823cff,
  transparent
 );
 margin:20px 0;
}


/* =========================================================
   BOTÕES
   ========================================================= */

.stButton button,
.stDownloadButton button{
 border-radius:10px!important;
 border:1px solid rgba(68,137,255,.55)!important;
 background:linear-gradient(
  100deg,
  #0876df,
  #5631d6
 )!important;
 color:#fff!important;
 font-weight:700;
 min-height:42px;
 transition:.2s;
}

.stButton button:hover,
.stDownloadButton button:hover{
 transform:translateY(-2px);
}


/*
 Espaço utilizado para colocar os botões exatamente
 na altura dos campos de seleção.
*/
.filtro-botao{
 height:28px;
 width:100%;
}


/* Remove espaçamentos extras do botão de download */
div[data-testid="stDownloadButton"]{
 margin-top:0!important;
}


/* =========================================================
   SELECTBOX
   ========================================================= */

div[data-baseweb="select"]>div{
 background:#0c111e!important;
 border-color:rgba(76,122,255,.3)!important;
}

div[data-baseweb="select"] span{
 color:#e8efff!important;
}

.stSelectbox label{
 color:#8cdfff!important;
 font-weight:700!important;
}


/* =========================================================
   CARDS
   ========================================================= */

.card{
 padding:18px;
 border-radius:15px;
 border:1px solid rgba(87,125,210,.22);
 background:linear-gradient(
  145deg,
  #0e1320,
  #070b14
 );
 min-height:110px;
}

.card-label{
 color:#9ca8bc;
 font-size:.78rem;
 font-weight:700;
}

.card-value{
 font-family:Orbitron;
 font-size:1.55rem;
 font-weight:700;
 margin-top:8px;
}

.card-foot{
 color:#718097;
 font-size:.7rem;
 margin-top:5px;
}


/* =========================================================
   PAINÉIS
   ========================================================= */

.panel{
 border:1px solid rgba(88,130,230,.2);
 background:linear-gradient(
  145deg,
  #0b101c,
  #060a12
 );
 border-radius:15px;
 padding:16px;
}

.panel-title{
 font-size:1rem;
 font-weight:800;
 margin-bottom:5px;
}

.panel-sub{
 color:#77849a;
 font-size:.75rem;
}


/* =========================================================
   TÍTULOS
   ========================================================= */

.section-title{
 font-size:1.2rem;
 font-weight:800;
 margin:8px 0;
}

.section-sub{
 color:#7d8ba2;
 font-size:.78rem;
}


/* =========================================================
   TABELA
   ========================================================= */

.table-wrap{
 width:100%;
 max-height:500px;
 overflow:auto;
 border:1px solid rgba(76,122,255,.25);
 border-radius:12px;
 background:#08101c;
}

.amira-table{
 width:100%;
 border-collapse:collapse;
 color:#e9f1ff;
 font-size:.82rem;
}

.amira-table th{
 position:sticky;
 top:0;
 z-index:2;
 padding:11px;
 text-align:left;
 color:#fff;
 background:#101b2d;
 border-bottom:1px solid #3564a5;
 white-space:nowrap;
}

.amira-table td{
 padding:9px 11px;
 border-bottom:1px solid rgba(76,122,255,.12);
 background:#0a1422;
 white-space:nowrap;
}

.amira-table tr:nth-child(even) td{
 background:#0c1727;
}

.amira-table tr:hover td{
 background:#10223a;
}

.num{
 text-align:right;
}


/* =========================================================
   FILTRO
   ========================================================= */

.filter{
 margin:10px 0;
 padding:11px 14px;
 border-radius:10px;
 border:1px solid rgba(32,232,137,.35);
 background:rgba(8,61,42,.45);
 color:#bfe9d3;
 font-size:.78rem;
}

.filter b{
 color:#37ef9b;
}


/* =========================================================
   RODAPÉ
   ========================================================= */

.footer{
 text-align:center;
 color:#59677d;
 font-size:.68rem;
 padding:20px 0;
}

</style>
""", unsafe_allow_html=True)


# =========================================================
# FUNÇÕES
# =========================================================

def vazio():
    return pd.DataFrame(
        columns=[
            "Data",
            "Hora",
            "Peso (kg)",
            "Lote"
        ]
    )


def data_planilha(v):
    if pd.isna(v) or str(v).strip() == "":
        return None

    s = str(v).strip()

    # Se contiver a letra 'T' (formato ISO/Sheets), pega só a parte da data YYYY-MM-DD ou extrai via Regex
    if "T" in s:
        s = s.split("T")[0]

    # Tenta converter no formato DD/MM/YYYY
    try:
        d = pd.to_datetime(s, format="%d/%m/%Y", errors="coerce")
        if not pd.isna(d):
            return d.strftime("%d/%m/%Y")
    except Exception:
        pass

    # Tenta converter qualquer outro formato sem considerar UTC/Timezone
    try:
        d = pd.to_datetime(s, errors="coerce")
        if not pd.isna(d):
            return d.strftime("%d/%m/%Y")
    except Exception:
        pass

    return s


def hora_planilha(v):
    if pd.isna(v) or str(v).strip() == "":
        return None

    s = str(v).strip()

    # 1. Se vier no formato ISO do Sheets (ex: 1899-12-30T20:54:13.000Z),
    # corrige o deslocamento de 8 horas antes de exibir a hora
    m_iso = re.search(r"T(\d{2}:\d{2}:\d{2})", s)
    if m_iso:
        h, m, sec = map(int, m_iso.group(1).split(":"))
        total_segundos = (h * 3600 + m * 60 + sec - 8 * 3600) % (24 * 3600)
        h = total_segundos // 3600
        m = (total_segundos % 3600) // 60
        sec = total_segundos % 60
        return f"{h:02d}:{m:02d}:{sec:02d}"

    # 2. Formato comum HH:MM:SS ou HH:MM
    m = re.match(r"^(\d{1,2}):(\d{2})(?::(\d{2}))?$", s)
    if m:
        return f"{int(m.group(1)):02d}:{m.group(2)}:{m.group(3) or '00'}"

    # 3. Se contiver a letra 'Z' ou offset de fuso, remove antes de converter
    s_limpo = re.sub(r"(Z|[+-]\d{2}:\d{2})$", "", s)

    try:
        d = pd.to_datetime(s_limpo, errors="coerce")
        if not pd.isna(d):
            return d.strftime("%H:%M:%S")
    except Exception:
        pass

    return s
def lote(v):

    if pd.isna(v) or str(v).strip() == "":
        return "Sem lote"

    s = re.sub(
        r"\s+",
        " ",
        str(v).strip()
    )

    if re.fullmatch(r"\d+\.0+", s):
        s = s.split(".")[0]

    return s


def normalizar(df):

    if df is None or df.empty:
        return vazio()

    df = df.copy()
    nomes = {}

    for c in df.columns:

        k = (
            str(c)
            .strip()
            .lower()
            .replace(" ", "")
            .replace("_", "")
        )

        if k in [
            "peso",
            "peso(kg)",
            "pesokg"
        ]:
            nomes[c] = "Peso (kg)"

        elif k == "data":
            nomes[c] = "Data"

        elif k == "hora":
            nomes[c] = "Hora"

        elif k == "lote":
            nomes[c] = "Lote"

    df.rename(
        columns=nomes,
        inplace=True
    )

    for c in [
        "Data",
        "Hora",
        "Peso (kg)",
        "Lote"
    ]:
        if c not in df:
            df[c] = None

    df["Peso (kg)"] = pd.to_numeric(
        df["Peso (kg)"],
        errors="coerce"
    )

    df["Data"] = df["Data"].map(
        data_planilha
    )

    df["Hora"] = df["Hora"].map(
        hora_planilha
    )

    df["Lote"] = df["Lote"].map(
        lote
    )

    return df[
        [
            "Data",
            "Hora",
            "Peso (kg)",
            "Lote"
        ]
    ].dropna(
        subset=["Peso (kg)"],
        how="all"
    ).reset_index(drop=True)


@st.cache_data(ttl=15)
def carregar():

    r = requests.get(
        URL_SCRIPT,
        timeout=20
    )

    r.raise_for_status()

    dados = r.json()

    if not isinstance(dados, dict):
        raise ValueError(
            "O Apps Script não devolveu as abas corretamente."
        )

    saida = {}

    for nome, conteudo in dados.items():

        if (
            not isinstance(conteudo, list)
            or len(conteudo) < 1
        ):
            continue

        cab = conteudo[0]
        linhas = conteudo[1:]

        if isinstance(cab, list):

            saida[str(nome)] = normalizar(
                pd.DataFrame(
                    linhas,
                    columns=cab
                )
            )

    return saida


def data_aba(nome):

    s = str(nome)

    for padrao, formato in [

        (
            r"(?<!\d)(\d{2}-\d{2}-\d{4})(?!\d)",
            "%d-%m-%Y"
        ),

        (
            r"(?<!\d)(\d{4}-\d{2}-\d{2})(?!\d)",
            "%Y-%m-%d"
        )

    ]:

        m = re.search(
            padrao,
            s
        )

        if m:

            try:
                return datetime.strptime(
                    m.group(1),
                    formato
                ).date()

            except Exception:
                pass

    return None


def rotulo(nome):

    d = data_aba(nome)

    return (
        d.strftime("%d/%m/%Y")
        if d
        else str(nome)
    )


def chave(nome):

    d = data_aba(nome)

    return (
        datetime.combine(
            d,
            datetime.min.time()
        )
        if d
        else datetime.min
    )


def numero(v):

    return (
        f"{v:,.2f}"
        .replace(",", "X")
        .replace(".", ",")
        .replace("X", ".")
    )


def tabela(df, altura=500):

    if df is None or df.empty:

        st.info(
            "Nenhum registro para exibir."
        )

        return

    h = f"""
    <div class="table-wrap"
         style="max-height:{altura}px">

    <table class="amira-table">

    <thead>
    <tr>
    """

    for c in df.columns:
        h += f"<th>{c}</th>"

    h += """
    </tr>
    </thead>
    <tbody>
    """

    for _, row in df.iterrows():

        h += "<tr>"

        for c in df.columns:

            v = row[c]

            if pd.isna(v):
                v = ""

            elif c in [
                "Peso (kg)",
                "Peso total (kg)",
                "Média (kg)"
            ]:

                try:
                    v = numero(float(v))
                except Exception:
                    pass

            h += f"<td>{v}</td>"

        h += "</tr>"

    h += """
    </tbody>
    </table>
    </div>
    """

    st.markdown(
        h,
        unsafe_allow_html=True
    )


def layout(fig, altura=300):

    fig.update_layout(
        template="plotly_dark",
        paper_bgcolor="rgba(0,0,0,0)",
        plot_bgcolor="rgba(0,0,0,0)",
        font=dict(
            color="#aeb9cb",
            family="Inter"
        ),
        margin=dict(
            l=10,
            r=10,
            t=20,
            b=10
        ),
        height=altura,
        showlegend=False
    )

    fig.update_xaxes(
        gridcolor="rgba(90,120,180,.10)",
        zeroline=False
    )

    fig.update_yaxes(
        gridcolor="rgba(90,120,180,.10)",
        zeroline=False
    )

    return fig


def agrupar_lote(df):

    if df.empty:
        return pd.DataFrame()

    x = df.copy()

    x["Lote"] = x["Lote"].map(lote)

    return x.groupby(
        "Lote"
    ).agg(
        Pesagens=("Peso (kg)", "count"),
        **{
            "Peso total (kg)": (
                "Peso (kg)",
                "sum"
            ),
            "Média (kg)": (
                "Peso (kg)",
                "mean"
            )
        }
    ).reset_index()


def metricas(df):

    caixas = len(df)

    peso = (
        float(df["Peso (kg)"].sum())
        if not df.empty
        else 0
    )

    media = (
        float(df["Peso (kg)"].mean())
        if not df.empty
        else 0
    )

    dados = [
        (
            "📦",
            "Caixas Passadas",
            caixas,
            "Registros filtrados"
        ),
        (
            "⚖",
            "Peso Total",
            f"{numero(peso)} kg",
            "Peso acumulado"
        ),
        (
            "◈",
            "Média por Caixa",
            f"{numero(media)} kg",
            "Média dos registros"
        )
    ]

    cols = st.columns(3)

    for col, (
        ico,
        nome,
        valor,
        rodape
    ) in zip(cols, dados):

        with col:

            st.markdown(
                f"""<div class="card">
<div class="card-label">{ico} &nbsp; {nome}</div>
<div class="card-value">{valor}</div>
<div class="card-foot">{rodape}</div>
</div>""",
                unsafe_allow_html=True
            )


# =========================================================
# EXPORTAÇÃO EXCEL
# =========================================================

def xlsx(df):

    b = io.BytesIO()

    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill

    wb = Workbook()

    ws = wb.active

    ws.title = "Registros"

    ws.append(
        list(df.columns)
    )

    for c in ws[1]:

        c.font = Font(
            bold=True,
            color="FFFFFF"
        )

        c.fill = PatternFill(
            "solid",
            fgColor="111827"
        )

    for row in df.itertuples(
        index=False,
        name=None
    ):

        ws.append(
            list(row)
        )

    ws.freeze_panes = "A2"

    ws.auto_filter.ref = (
        ws.dimensions
    )

    wb.save(b)

    return b.getvalue()


# =========================================================
# CARREGAR DADOS
# =========================================================

try:

    dados = carregar()
    erro = None

except Exception as e:

    dados = {}
    erro = str(e)


abas = sorted(
    dados.keys(),
    key=chave,
    reverse=True
)

if erro:

    st.error(
        f"Erro ao carregar os dados: {erro}"
    )


# =========================================================
# SIDEBAR
# =========================================================

with st.sidebar:

    if os.path.exists("logo.png"):

        st.image(
            "logo.png",
            use_container_width=True
        )

    elif os.path.exists("logo.jpg"):

        st.image(
            "logo.jpg",
            use_container_width=True
        )

    st.markdown(
        '<div class="brand">'
        'Sistema <span>AMIRA</span>'
        '</div>'
        '<div class="caption">'
        'Monitoramento • Automação • Precisão'
        '</div>',
        unsafe_allow_html=True
    )

    st.markdown("### MENU")

    menus = [
        "📋  Registro e Dados",
        "📊  Gráficos e Análises",
        "🗂  Histórico de Planilhas"
    ]

    menu = st.session_state.get(
        "menu",
        menus[0]
    )

    for i, m in enumerate(menus):

        if st.button(
            m,
            key=f"menu_{i}",
            type=(
                "primary"
                if menu == m
                else "secondary"
            ),
            use_container_width=True
        ):

            st.session_state.menu = m
            st.rerun()

    st.markdown(
        '<br>'
        '<div class="panel">'
        '🟢 <b>Sistema AMIRA</b><br>'
        '<span style="color:#7f8da4;'
        'font-size:.75rem">'
        'Monitoramento ativo'
        '</span>'
        '</div>',
        unsafe_allow_html=True
    )


# =========================================================
# CABEÇALHO
# =========================================================

titulos = {

    menus[0]: (
        "Bem-vindo à <span>AMIRA</span>",
        "Sistema de Monitoramento e Registro de Produção"
    ),

    menus[1]: (
        "<span>Gráficos</span> e Análises",
        "Explore o histórico e compare a produção."
    ),

    menus[2]: (
        "<span>Histórico</span> de Planilhas",
        "Controle dos lançamentos registrados."
    )

}

titulo, subtitulo = titulos[menu]

st.markdown(
    f'<div class="top-title">{titulo}</div>'
    f'<div class="subtitle">{subtitulo}</div>'
    '<div class="line"></div>',
    unsafe_allow_html=True
)


# =========================================================
# FILTROS E AÇÕES
# =========================================================

dia = st.session_state.get(
    "dia",
    abas[0] if abas else None
)

lote_atual = st.session_state.get(
    "lote",
    "Todos os lotes"
)

c1, c2, c3, c4 = st.columns(
    [1.3, 1.3, .9, .9]
)


# =========================================================
# DIA DE PRODUÇÃO
# =========================================================

with c1:

    if menu == menus[0] and abas:

        novo_dia = st.selectbox(
            "📅 Dia de Produção",
            abas,
            index=(
                abas.index(dia)
                if dia in abas
                else 0
            ),
            format_func=rotulo
        )

        if novo_dia != dia:

            st.session_state.dia = novo_dia
            st.session_state.lote = "Todos os lotes"

            st.rerun()

        dia = novo_dia


# =========================================================
# DATA SELECIONADA
# =========================================================

df = (
    dados.get(dia, vazio())
    if dia
    else vazio()
)


# =========================================================
# FILTRAR POR LOTE
# =========================================================

with c2:

    if menu == menus[0] and not df.empty:

        lotes = sorted(
            df["Lote"].map(lote).unique(),
            key=lambda x: (
                (0, int(x))
                if str(x).isdigit()
                else (1, str(x).lower())
            )
        )

        opcoes = [
            "Todos os lotes"
        ] + list(lotes)

        if lote_atual not in opcoes:
            lote_atual = "Todos os lotes"

        novo_lote = st.selectbox(
            "🔎 Filtrar por lote",
            opcoes,
            index=opcoes.index(
                lote_atual
            )
        )

        if novo_lote != lote_atual:

            st.session_state.lote = novo_lote
            st.rerun()

        lote_atual = novo_lote


# =========================================================
# APLICA FILTRO
# =========================================================

df_filtro = (
    df[
        df["Lote"].map(lote)
        == lote_atual
    ].copy()
    if lote_atual != "Todos os lotes"
    else df.copy()
)


# =========================================================
# RECARREGAR
# =========================================================

with c3:

    if menu == menus[0]:

        st.markdown(
            '<div class="filtro-botao"></div>',
            unsafe_allow_html=True
        )

        if st.button(
            "🔄 Recarregar",
            key="btn_recarregar",
            use_container_width=True
        ):

            carregar.clear()

            st.rerun()


# =========================================================
# BAIXAR PLANILHA
# =========================================================

with c4:

    if menu == menus[0] and not df_filtro.empty:

        st.markdown(
            '<div class="filtro-botao"></div>',
            unsafe_allow_html=True
        )

        st.download_button(
            "⬇️ Baixar planilha",
            data=xlsx(df_filtro),
            file_name=(
                f"AMIRA_"
                f"{rotulo(dia).replace('/', '-')}"
                f".xlsx"
            ),
            mime=(
                "application/vnd.openxmlformats-"
                "officedocument.spreadsheetml.sheet"
            ),
            key="btn_baixar",
            use_container_width=True
        )


# =========================================================
# REGISTRO E DADOS
# =========================================================

if menu == menus[0]:

    if lote_atual != "Todos os lotes":

        st.markdown(
            f'<div class="filter">'
            f'🔎 Filtro ativo: '
            f'<b>Lote {lote_atual}</b>'
            f'</div>',
            unsafe_allow_html=True
        )

    metricas(df_filtro)

    st.markdown("<br>", unsafe_allow_html=True)


    # =====================================================
    # PRODUÇÃO POR LOTE
    # =====================================================

    st.markdown(
        '<div class="section-title">'
        '📦 Produção por Lote'
        '</div>'
        '<div class="section-sub">'
        'Agrupamento automático das pesagens.'
        '</div>',
        unsafe_allow_html=True
    )

    agrupado = agrupar_lote(
        df_filtro
    )

    if not agrupado.empty:

        a, b = st.columns(
            [1, 1.2]
        )

        with a:

            tabela(
                agrupado,
                350
            )

        with b:

            fig = px.bar(
                agrupado,
                x="Lote",
                y="Peso total (kg)",
                text="Peso total (kg)"
            )

            fig.update_traces(
                marker_color="#2f80ff",
                texttemplate="%{text:.2f}",
                textposition="outside"
            )

            st.plotly_chart(
                layout(fig, 280),
                use_container_width=True,
                config={
                    "displayModeBar": False
                }
            )

    else:

        st.info(
            "Nenhum lote disponível."
        )


    # =====================================================
    # REGISTROS
    # =====================================================

    st.markdown(
        '<div class="section-title">'
        '📋 Registros de Produção'
        '</div>'
        '<div class="section-sub">'
        'Pesagens do dia selecionado.'
        '</div>',
        unsafe_allow_html=True
    )

    if df_filtro.empty:

        st.info(
            "Aguardando pesagens para "
            "exibir na tabela."
        )

    else:

        tabela_dia = (
            df_filtro
            .copy()
            .reset_index(drop=True)
        )

        tabela_dia.insert(
            0,
            "#",
            tabela_dia.index + 1
        )

        tabela(
            tabela_dia,
            440
        )


    # =====================================================
    # GRÁFICOS DO DIA
    # =====================================================

    st.markdown(
        '<br>'
        '<div class="section-title">'
        '📈 Desempenho do Dia'
        '</div>',
        unsafe_allow_html=True
    )

    g1, g2 = st.columns(2)

    if not df_filtro.empty:

        temp = (
            df_filtro
            .reset_index(drop=True)
            .copy()
        )

        temp["Registro"] = (
            temp.index + 1
        )

        with g1:

            st.markdown(
                "**Peso ao longo do dia**"
            )

            fig = px.area(
                temp,
                x="Registro",
                y="Peso (kg)"
            )

            fig.update_traces(
                line_color="#3d8cff",
                fillcolor="rgba(61,140,255,.15)"
            )

            st.plotly_chart(
                layout(fig),
                use_container_width=True,
                config={
                    "displayModeBar": False
                }
            )

        with g2:

            st.markdown(
                "**Distribuição dos pesos**"
            )

            distribuicao = (
                df_filtro["Peso (kg)"]
                .round(2)
                .value_counts()
                .sort_index()
                .rename_axis("Peso (kg)")
                .reset_index(name="Quantidade")
            )

            fig = px.bar(
                distribuicao,
                x="Peso (kg)",
                y="Quantidade"
            )

            fig.update_traces(
                marker_color="#8a35ff",
                text="Quantidade",
                textposition="outside"
            )

            fig.update_xaxes(
                title="Peso (kg)",
                type="category"
            )

            fig.update_yaxes(
                title="Quantidade de registros",
                dtick=1
            )

            st.plotly_chart(
                layout(fig),
                use_container_width=True,
                config={
                    "displayModeBar": False
                }
            )


# =========================================================
# GRÁFICOS E ANÁLISES
# =========================================================

elif menu == menus[1]:

    resumo = []

    for nome, temp in dados.items():

        if not temp.empty:

            resumo.append(
                {
                    "Data": rotulo(nome),
                    "Peso total (kg)": (
                        temp["Peso (kg)"].sum()
                    ),
                    "Caixas": len(temp)
                }
            )

    rdf = pd.DataFrame(resumo)

    if not rdf.empty:
        rdf["Data"] = rdf["Data"].str[:5]

    if rdf.empty:

        st.info(
            "Aguardando dados para gerar os gráficos."
        )

    else:

        st.markdown(
            '<div class="section-title">'
            '📊 Comparativo Diário'
            '</div>',
            unsafe_allow_html=True
        )

        c1, c2 = st.columns(2)

        with c1:

            fig = px.bar(
                rdf,
                x="Data",
                y="Peso total (kg)"
            )

            fig.update_traces(
                marker_color="#2f80ff"
            )

            st.plotly_chart(
                layout(fig),
                use_container_width=True,
                config={
                    "displayModeBar": False
                }
            )

        with c2:

            fig = px.bar(
                rdf,
                x="Data",
                y="Caixas"
            )

            fig.update_traces(
                marker_color="#8a35ff"
            )

            st.plotly_chart(
                layout(fig),
                use_container_width=True,
                config={
                    "displayModeBar": False
                }
            )


        # =================================================
        # ANÁLISE MENSAL
        # =================================================

        mensal = []

        for nome, temp in dados.items():

            d = data_aba(nome)

            if d and not temp.empty:

                mensal.append(
                    {
                        "Ano": d.year,
                        "Mês": d.strftime("%m/%Y"),
                        "Mês_Num": d.month,
                        "Peso": (
                            temp["Peso (kg)"].sum()
                        ),
                        "Caixas": len(temp)
                    }
                )

        if mensal:

            m = pd.DataFrame(
                mensal
            )

            m = (
                m.groupby(
                    [
                        "Ano",
                        "Mês",
                        "Mês_Num"
                    ]
                )[["Peso", "Caixas"]]
                .sum()
                .reset_index()
                .sort_values(
                    ["Ano", "Mês_Num"]
                )
            )

            st.markdown(
                '<br>'
                '<div class="section-title">'
                '📅 Análise Mensal'
                '</div>',
                unsafe_allow_html=True
            )

            c1, c2 = st.columns(2)

            with c1:

                fig = px.bar(
                    m,
                    x="Mês",
                    y="Peso",
                    color="Ano",
                    color_discrete_sequence=[
                        "#00d2ff",
                        "#2f80ff",
                        "#8a35ff"
                    ]
                )

                st.plotly_chart(
                    layout(fig),
                    use_container_width=True,
                    config={
                        "displayModeBar": False
                    }
                )

            with c2:

                fig = px.line(
                    m,
                    x="Mês",
                    y="Caixas",
                    color="Ano",
                    markers=True,
                    color_discrete_sequence=[
                        "#00d2ff",
                        "#2f80ff",
                        "#8a35ff"
                    ]
                )

                fig.update_traces(
                    line_width=3
                )

                st.plotly_chart(
                    layout(fig),
                    use_container_width=True,
                    config={
                        "displayModeBar": False
                    }
                )


# =========================================================
# HISTÓRICO
# =========================================================

else:

    if not abas:

        st.info(
            "Ainda não há planilhas registradas."
        )

    else:

        # =====================================================
        # PLANILHAS / DADOS DO DIA
        # A planilha continua disponível para consulta.
        # O botão de envio fica dentro de cada dia para não
        # deixar o histórico visualmente carregado.
        # =====================================================

        for nome in abas:

            temp = dados.get(
                nome,
                vazio()
            )

            rot = rotulo(nome)
            chave_envio = f"historico_enviado_{nome}"

            if chave_envio not in st.session_state:
                st.session_state[chave_envio] = False

            with st.expander(
                f"📄 {rot} — {'✅ Enviado' if st.session_state[chave_envio] else '🟡 Pendente'}"
            ):

                if st.session_state[chave_envio]:
                    if st.button(
                        "↩️ Marcar como pendente",
                        key=f"marcar_pendente_{nome}",
                        use_container_width=False
                    ):
                        st.session_state[chave_envio] = False
                        st.rerun()
                else:
                    if st.button(
                        "📤 Marcar como enviada",
                        key=f"marcar_enviada_{nome}",
                        use_container_width=False
                    ):
                        st.session_state[chave_envio] = True
                        st.rerun()

                if temp.empty:

                    st.info(
                        "A planilha ainda não possui registros."
                    )

                else:

                    hist = (
                        temp
                        .copy()
                        .reset_index(drop=True)
                    )

                    hist.insert(
                        0,
                        "#",
                        hist.index + 1
                    )

                    tabela(
                        hist,
                        400
                    )

                    resumo = agrupar_lote(
                        temp
                    )

                    if not resumo.empty:

                        st.markdown(
                            "#### 📦 Resumo por lote"
                        )

                        tabela(
                            resumo,
                            300
                        )

# RODAPÉ
# =========================================================

st.markdown(
    '<div class="footer">'
    'AMIRA • Sistema de Monitoramento '
    'e Registro de Produção • SENAI'
    '</div>',
    unsafe_allow_html=True
)
